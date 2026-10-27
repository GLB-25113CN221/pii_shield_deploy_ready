"""DOCX redaction across paragraphs, tables, headers and footers."""
from io import BytesIO
from docx import Document
from detectors.base import Finding
from redactors.text import redact_text
def redact_docx(data: bytes, detector) -> bytes:
    doc=Document(BytesIO(data))
    def paragraph(p):
        original=p.text
        findings=detector(original)
        if findings:
            # Replacing paragraph text guarantees removal, but may simplify formatting.
            p.text=redact_text(original,findings)
    def table(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs: paragraph(p)
                for nested in cell.tables: table(nested)
    for p in doc.paragraphs: paragraph(p)
    for t in doc.tables: table(t)
    for section in doc.sections:
        for p in section.header.paragraphs: paragraph(p)
        for t in section.header.tables: table(t)
        for p in section.footer.paragraphs: paragraph(p)
        for t in section.footer.tables: table(t)
    out=BytesIO(); doc.core_properties.author=""; doc.core_properties.title=""; doc.core_properties.subject=""
    doc.save(out); return out.getvalue()

def docx_text(data: bytes) -> str:
    """All paragraph/table text of a DOCX (used for the report and risk score)."""
    doc=Document(BytesIO(data)); parts=[p.text for p in doc.paragraphs]
    def walk(t):
        for row in t.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
                for n in cell.tables: walk(n)
    for t in doc.tables: walk(t)
    return "\n".join(parts)
