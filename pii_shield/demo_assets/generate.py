"""Generate synthetic demonstration assets only; never use real identifiers."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
try:
    import pymupdf as fitz
except ImportError:
    import fitz
def generate(folder=None):
    p=Path(folder) if folder else Path(__file__).parent/"generated"; p.mkdir(parents=True,exist_ok=True)
    (p/"sample.txt").write_text("नाम: राजेश कुमार\nEmail: demo.user@example.com\nPAN: ABCPE1234F\n",encoding="utf-8")
    im=Image.new("RGB",(900,500),"white"); d=ImageDraw.Draw(im)
    try: font=ImageFont.load_default(size=32)   # larger text = better OCR
    except TypeError: font=ImageFont.load_default()
    d.text((40,40),"DEMO ID CARD — SYNTHETIC DATA",fill="black",font=font)
    d.text((40,120),"नाम: राजेश कुमार",fill="black",font=font); d.text((40,180),"AADHAAR (FAKE): 234567890124",fill="black",font=font)
    im.save(p/"sample_id.png")
    doc=Document(); doc.add_heading("Synthetic sample",0); doc.add_paragraph("नाम: प्रिया शर्मा"); doc.add_paragraph("Email: demo.user@example.com"); doc.save(p/"sample.docx")
    pdf=fitz.open(); page=pdf.new_page(); page.insert_text((50,70),"Synthetic sample PDF")
    page.insert_text((50,110),"Email: demo.user@example.com")
    pdf.save(str(p/"sample.pdf"))
    pix=page.get_pixmap(matrix=fitz.Matrix(2,2)); scan=fitz.open(); pg=scan.new_page(width=pix.width,height=pix.height); pg.insert_image(pg.rect,stream=pix.tobytes("png")); scan.save(str(p/"sample_scanned.pdf")); pdf.close(); scan.close()
    return [str(x) for x in p.iterdir()]
if __name__=="__main__": print("\n".join(generate()))
